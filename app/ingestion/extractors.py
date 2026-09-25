"""
Extracts raw text from an uploaded file.

Named edge cases this module explicitly handles (not just the happy path):
  1. Empty documents            -> EmptyDocumentError
  2. Malformed / unparseable    -> ExtractionError
  3. Non-English text           -> allowed through, but flagged as a warning
  4. Scanned PDFs with no text  -> treated as empty (we do not OCR, and say so)
"""
from dataclasses import dataclass, field
from pathlib import Path
from typing import List

from pypdf import PdfReader
from pypdf.errors import PdfReadError
import docx as python_docx

from app.exceptions import EmptyDocumentError, ExtractionError, UnsupportedFileTypeError

try:
    from langdetect import detect, LangDetectException
    _LANGDETECT_AVAILABLE = True
except ImportError:  # pragma: no cover - defensive; langdetect is a listed dependency
    _LANGDETECT_AVAILABLE = False


@dataclass
class ExtractedDocument:
    text: str
    warnings: List[str] = field(default_factory=list)


def _detect_language_warning(text: str) -> List[str]:
    """Best-effort language check. Never blocks ingestion -- only warns,
    because rejecting non-English content outright would be a product
    decision, not an engineering one. The retrieval/embedding model in the
    default config is English-optimized, so we're honest about that."""
    if not _LANGDETECT_AVAILABLE:
        return []
    sample = text[:2000].strip()
    if len(sample) < 40:
        return []  # too short for a reliable language guess
    try:
        lang = detect(sample)
    except LangDetectException:
        return []
    if lang != "en":
        return [
            f"Detected non-English content (lang='{lang}'). The default embedding "
            "model is English-optimized; retrieval quality may be degraded."
        ]
    return []


def _extract_pdf(path: Path) -> str:
    try:
        reader = PdfReader(str(path))
    except PdfReadError as e:
        raise ExtractionError(f"Could not parse PDF '{path.name}': {e}") from e
    except Exception as e:  # noqa: BLE001 - we want to convert ANY parse failure
        raise ExtractionError(f"Unexpected error reading PDF '{path.name}': {e}") from e

    if reader.is_encrypted:
        raise ExtractionError(f"PDF '{path.name}' is password-protected; cannot extract text.")

    pages_text = []
    for page in reader.pages:
        try:
            pages_text.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - a single bad page shouldn't kill the whole doc
            pages_text.append("")
    return "\n\n".join(pages_text)


def _extract_docx(path: Path) -> str:
    try:
        document = python_docx.Document(str(path))
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"Could not parse DOCX '{path.name}': {e}") from e
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_plain(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"Could not read text file '{path.name}': {e}") from e


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".txt": _extract_plain,
    ".md": _extract_plain,
}


def extract_text(path: Path, supported_extensions: List[str]) -> ExtractedDocument:
    ext = path.suffix.lower()
    if ext not in supported_extensions:
        raise UnsupportedFileTypeError(
            f"'{ext}' is not supported. Supported types: {', '.join(supported_extensions)}"
        )

    raw_text = _EXTRACTORS[ext](path)
    cleaned = raw_text.strip()

    if not cleaned:
        raise EmptyDocumentError(
            f"'{path.name}' produced no extractable text. If this is a scanned "
            "PDF (image-only), this pipeline does not run OCR -- run it through "
            "an OCR step first."
        )

    warnings = _detect_language_warning(cleaned)
    return ExtractedDocument(text=cleaned, warnings=warnings)
