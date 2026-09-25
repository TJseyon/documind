"""
One-time generator for the sample .docx and .pdf files under sample_docs/.
Run with: python scripts/generate_sample_docs.py

These exist purely so the eval harness has something to point at out of the
box. They are NOT a substitute for testing against your own real documents.
"""
from pathlib import Path

from docx import Document
from fpdf import FPDF

OUT_DIR = Path(__file__).resolve().parent.parent / "sample_docs"


def make_employee_handbook_docx():
    doc = Document()
    doc.add_heading("Employee Handbook (Sample)", level=1)

    doc.add_heading("Paid Time Off", level=2)
    doc.add_paragraph(
        "Full-time employees receive 20 vacation days per calendar year, "
        "accrued monthly. Unused vacation days roll over up to a maximum of "
        "5 days into the following year."
    )
    doc.add_paragraph(
        "Full-time employees also receive 10 sick days per calendar year. "
        "Sick days do not roll over and do not require advance notice."
    )

    doc.add_heading("Benefits Enrollment", level=2)
    doc.add_paragraph(
        "New employees can enroll in health insurance and the 401k "
        "retirement plan during their first 30 days. Contact HR to begin "
        "enrollment; HR manages all benefits paperwork and vendor "
        "coordination."
    )

    doc.add_heading("Remote Work Policy", level=2)
    doc.add_paragraph(
        "Employees may work remotely up to 3 days per week with manager "
        "approval. Fully remote arrangements require VP-level sign-off."
    )

    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Benefit", "Amount"
    row = table.add_row().cells
    row[0].text, row[1].text = "Annual vacation days", "20"
    row = table.add_row().cells
    row[0].text, row[1].text = "Annual sick days", "10"

    doc.save(OUT_DIR / "employee_handbook.docx")


def make_product_spec_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "SR-40 Drone - Product Specification (Sample)", ln=True)
    pdf.set_font("Helvetica", size=11)
    pdf.ln(4)

    body = (
        "The SR-40 is a mid-range commercial inspection drone.\n\n"
        "Payload: The SR-40 has a maximum payload capacity of 2.5 kilograms, "
        "sufficient for standard inspection camera and sensor packages.\n\n"
        "Flight time: On a single charge, the SR-40 achieves 28 minutes of "
        "flight time under standard load, or approximately 22 minutes at "
        "maximum payload.\n\n"
        "Altitude: The SR-40 has a maximum operating altitude of 120 meters "
        "above ground level, in line with common commercial drone "
        "regulations.\n\n"
        "Range: Maximum control range is 6 kilometers in open, "
        "unobstructed conditions.\n\n"
        "Charging: A full charge from empty takes approximately 65 minutes "
        "using the included fast charger."
    )
    pdf.multi_cell(0, 7, body)
    pdf.output(str(OUT_DIR / "product_spec.pdf"))


def make_edge_case_docs():
    (OUT_DIR / "empty_document.txt").write_text("   \n\n   \n")
    (OUT_DIR / "non_english_sample.txt").write_text(
        "Este es un documento de ejemplo escrito completamente en espa\u00f1ol. "
        "Contiene varias oraciones para que el detector de idioma tenga "
        "suficiente texto para analizar con confianza."
    )


if __name__ == "__main__":
    OUT_DIR.mkdir(exist_ok=True)
    make_employee_handbook_docx()
    make_product_spec_pdf()
    make_edge_case_docs()
    print(f"Sample docs written to {OUT_DIR}")
